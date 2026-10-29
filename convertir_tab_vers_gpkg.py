# -*- coding: utf-8 -*-
"""
Conversion récursive des fichiers MapInfo TAB en GeoPackage.
- Ne crée/conserve un GeoPackage que si la conversion réussit et contient au moins une entité.
- Supprime le fichier temporaire en cas d'échec ou de couche vide.
- Génère un rapport Excel avec le résultat de chaque fichier TAB.

À exécuter avec le Python fourni avec QGIS (testé avec QGIS 3.40).

Configuration :
    Les chemins de dossier se définissent soit dans la section PARAMÈTRES ci-dessous,
    soit via les variables d'environnement TAB_SOURCE_DIR et GPKG_OUTPUT_DIR.

Exemple d'utilisation avec le Python de QGIS :
    "C:\Program Files\QGIS\apps\Python312\python.exe" convertir_tab_vers_gpkg.py
"""

import os
import traceback
from datetime import datetime

from osgeo import ogr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# PARAMÈTRES  (À ADAPTER AVANT EXÉCUTION)
# ---------------------------------------------------------------------------

# Dossier contenant les fichiers TAB à convertir (recherche récursive).
DOSSIER_SOURCE = os.environ.get(
    "TAB_SOURCE_DIR",
    r"CHEMIN/VERS/DOSSIER_SOURCE"  # <-- Remplacer par votre chemin source
)

# Dossier où seront écrits les GeoPackages et le rapport Excel.
NOM_DOSSIER_SORTIE = os.environ.get(
    "GPKG_OUTPUT_DIR",
    r"CHEMIN/VERS/DOSSIER_SORTIE"  # <-- Remplacer par votre chemin de sortie
)

NOM_RAPPORT = "rapport_conversion_mapinfo_vers_gpkg.xlsx"


# ---------------------------------------------------------------------------
# OUTILS
# ---------------------------------------------------------------------------

def supprimer_fichier_gpkg(path):
    """Supprime un GeoPackage et ses éventuels fichiers annexes."""
    for suffixe in ("", "-wal", "-shm", "-journal"):
        candidat = path + suffixe
        try:
            if os.path.isfile(candidat):
                os.remove(candidat)
        except OSError:
            pass


def compter_entites(couche):
    """Compte les entités, avec une méthode de secours si nécessaire."""
    try:
        nombre = couche.GetFeatureCount(1)
        if nombre >= 0:
            return nombre
    except Exception:
        pass

    nombre = 0
    couche.ResetReading()
    while couche.GetNextFeature() is not None:
        nombre += 1
    couche.ResetReading()
    return nombre


def copier_couche(source_couche, cible_ds, nom_couche):
    """Copie les champs et entités d'une couche OGR dans le GeoPackage."""
    srs = source_couche.GetSpatialRef()
    type_geom = source_couche.GetGeomType()

    cible_couche = cible_ds.CreateLayer(
        nom_couche,
        srs=srs,
        geom_type=type_geom
    )
    if cible_couche is None:
        raise RuntimeError(
            f"Impossible de créer la couche de sortie « {nom_couche} »."
        )

    # Copier la structure des champs.
    definition_source = source_couche.GetLayerDefn()
    for i in range(definition_source.GetFieldCount()):
        champ_source = definition_source.GetFieldDefn(i)
        champ_copie = ogr.FieldDefn(champ_source.GetNameRef(), champ_source.GetType())
        champ_copie.SetWidth(champ_source.GetWidth())
        champ_copie.SetPrecision(champ_source.GetPrecision())

        if cible_couche.CreateField(champ_copie) != ogr.OGRERR_NONE:
            raise RuntimeError(
                f"Impossible de créer le champ « {champ_source.GetNameRef()} »."
            )

    definition_cible = cible_couche.GetLayerDefn()
    source_couche.ResetReading()
    nombre_copie = 0

    for entite_source in source_couche:
        entite_cible = ogr.Feature(definition_cible)

        # Copier les valeurs attributaires par nom de champ.
        for i in range(definition_source.GetFieldCount()):
            nom_champ = definition_source.GetFieldDefn(i).GetNameRef()
            index_source = entite_source.GetFieldIndex(nom_champ)
            index_cible = entite_cible.GetFieldIndex(nom_champ)

            if index_source >= 0 and index_cible >= 0:
                if entite_source.IsFieldSetAndNotNull(index_source):
                    entite_cible.SetField(
                        index_cible,
                        entite_source.GetField(index_source)
                    )

        # Copier la géométrie, si elle existe.
        geometrie = entite_source.GetGeometryRef()
        if geometrie is not None:
            entite_cible.SetGeometry(geometrie.Clone())

        if cible_couche.CreateFeature(entite_cible) != ogr.OGRERR_NONE:
            entite_cible = None
            entite_source = None
            raise RuntimeError("Erreur pendant l'écriture d'une entité.")

        nombre_copie += 1
        entite_cible = None
        entite_source = None

    cible_couche.SyncToDisk()
    return nombre_copie


def convertir_tab(chemin_tab, dossier_sortie):
    """
    Convertit un fichier TAB vers un GeoPackage.
    Le GeoPackage final n'est publié que si la conversion est valide et non vide.
    Retourne un dictionnaire pour le rapport.
    """
    nom_base = os.path.splitext(os.path.basename(chemin_tab))[0]
    chemin_relatif = os.path.relpath(
        os.path.dirname(chemin_tab), DOSSIER_SOURCE
    )

    if chemin_relatif == ".":
        dossier_cible = dossier_sortie
    else:
        dossier_cible = os.path.join(dossier_sortie, chemin_relatif)

    os.makedirs(dossier_cible, exist_ok=True)

    chemin_gpkg = os.path.join(dossier_cible, nom_base + ".gpkg")
    chemin_temp = os.path.join(
        dossier_cible,
        f".{nom_base}_conversion_temporaire.gpkg"
    )

    resultat = {
        "Fichier TAB": chemin_tab,
        "GeoPackage": chemin_gpkg,
        "Statut": "",
        "Nombre d'entités": 0,
        "Message": ""
    }

    # Éliminer seulement un éventuel fichier temporaire laissé par un arrêt.
    supprimer_fichier_gpkg(chemin_temp)

    source_ds = None
    cible_ds = None

    try:
        source_ds = ogr.Open(chemin_tab, 0)
        if source_ds is None:
            raise RuntimeError("Impossible d'ouvrir le fichier MapInfo TAB.")

        if source_ds.GetLayerCount() == 0:
            resultat["Statut"] = "VIDE"
            resultat["Message"] = "Aucune couche lisible dans le fichier TAB."
            return resultat

        couche_source = source_ds.GetLayer(0)
        if couche_source is None:
            raise RuntimeError("Impossible de lire la couche du fichier TAB.")

        nombre_source = compter_entites(couche_source)
        if nombre_source == 0:
            resultat["Statut"] = "VIDE"
            resultat["Message"] = "La couche ne contient aucune entité. Aucun GPKG créé."
            return resultat

        pilote_gpkg = ogr.GetDriverByName("GPKG")
        if pilote_gpkg is None:
            raise RuntimeError("Le pilote OGR GeoPackage (GPKG) est indisponible.")

        cible_ds = pilote_gpkg.CreateDataSource(chemin_temp)
        if cible_ds is None:
            raise RuntimeError("Impossible de créer le GeoPackage temporaire.")

        nom_couche = nom_base[:60] or "couche"
        nombre_copie = copier_couche(couche_source, cible_ds, nom_couche)

        # Fermer les fichiers avant de vérifier et publier le résultat.
        cible_ds = None
        source_ds = None

        if nombre_copie == 0:
            supprimer_fichier_gpkg(chemin_temp)
            resultat["Statut"] = "VIDE"
            resultat["Message"] = "Aucune entité copiée. Aucun GPKG créé."
            return resultat

        # Vérifier que le GeoPackage temporaire s'ouvre et contient bien des entités.
        verification_ds = ogr.Open(chemin_temp, 0)
        if verification_ds is None or verification_ds.GetLayerCount() == 0:
            verification_ds = None
            raise RuntimeError("Vérification du GeoPackage temporaire impossible.")

        verification_couche = verification_ds.GetLayer(0)
        nombre_verifie = compter_entites(verification_couche)
        verification_ds = None

        if nombre_verifie == 0:
            supprimer_fichier_gpkg(chemin_temp)
            resultat["Statut"] = "VIDE"
            resultat["Message"] = "Le GeoPackage temporaire est vide. Aucun GPKG créé."
            return resultat

        # Publication uniquement après une conversion et une vérification réussies.
        if os.path.exists(chemin_gpkg):
            supprimer_fichier_gpkg(chemin_gpkg)
        os.replace(chemin_temp, chemin_gpkg)

        resultat["Statut"] = "SUCCÈS"
        resultat["Nombre d'entités"] = nombre_verifie
        resultat["Message"] = "Conversion réussie."
        return resultat

    except Exception as exc:
        resultat["Statut"] = "ÉCHEC"
        resultat["Message"] = f"{exc}"
        supprimer_fichier_gpkg(chemin_temp)
        return resultat

    finally:
        cible_ds = None
        source_ds = None
        # Suppression conditionnelle : après os.replace(), le temporaire n'existe plus.
        if os.path.exists(chemin_temp):
            supprimer_fichier_gpkg(chemin_temp)


def creer_rapport(resultats, chemin_rapport):
    """Crée un fichier Excel lisible contenant les résultats."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Rapport"

    colonnes = [
        "Fichier TAB",
        "GeoPackage",
        "Statut",
        "Nombre d'entités",
        "Message"
    ]
    feuille.append(colonnes)

    for resultat in resultats:
        feuille.append([resultat.get(colonne, "") for colonne in colonnes])

    couleur_entete = PatternFill("solid", fgColor="1F4E78")
    for cellule in feuille[1]:
        cellule.font = Font(color="FFFFFF", bold=True)
        cellule.fill = couleur_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")

    couleurs_statut = {
        "SUCCÈS": PatternFill("solid", fgColor="C6EFCE"),
        "VIDE": PatternFill("solid", fgColor="FFEB9C"),
        "ÉCHEC": PatternFill("solid", fgColor="FFC7CE"),
    }

    for ligne in range(2, feuille.max_row + 1):
        statut = feuille.cell(row=ligne, column=3).value
        if statut in couleurs_statut:
            feuille.cell(row=ligne, column=3).fill = couleurs_statut[statut]

    largeurs = {
        1: 75,
        2: 75,
        3: 16,
        4: 20,
        5: 75
    }
    for index, largeur in largeurs.items():
        feuille.column_dimensions[get_column_letter(index)].width = largeur

    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions

    classeur.save(chemin_rapport)


# ---------------------------------------------------------------------------
# PROGRAMME PRINCIPAL
# ---------------------------------------------------------------------------

def main():
    print("=" * 72)
    print("CONVERSION MAPINFO TAB VERS GEOPACKAGE")
    print("=" * 72)
    print(f"Dossier source : {DOSSIER_SOURCE}")

    if not os.path.isdir(DOSSIER_SOURCE):
        print("\nERREUR : le dossier source est introuvable ou inaccessible.")
        print("Vérifiez le chemin configuré (DOSSIER_SOURCE ou variable TAB_SOURCE_DIR).")
        input("\nAppuyez sur Entrée pour fermer...")
        return

    dossier_sortie = NOM_DOSSIER_SORTIE
    os.makedirs(dossier_sortie, exist_ok=True)

    fichiers_tab = []
    for racine, dossiers, fichiers in os.walk(DOSSIER_SOURCE):
        # Ne pas retraiter les sorties de précédentes exécutions.
        dossiers[:] = [
            dossier for dossier in dossiers
            if os.path.normcase(os.path.join(racine, dossier))
            != os.path.normcase(dossier_sortie)
        ]

        for nom_fichier in fichiers:
            if nom_fichier.lower().endswith(".tab"):
                fichiers_tab.append(os.path.join(racine, nom_fichier))

    fichiers_tab.sort(key=str.lower)

    if not fichiers_tab:
        print("\nAucun fichier .TAB trouvé.")
        input("\nAppuyez sur Entrée pour fermer...")
        return

    print(f"\n{len(fichiers_tab)} fichier(s) TAB trouvé(s).\n")

    resultats = []
    for numero, chemin_tab in enumerate(fichiers_tab, start=1):
        print(f"[{numero}/{len(fichiers_tab)}] {chemin_tab}")
        resultat = convertir_tab(chemin_tab, dossier_sortie)
        resultats.append(resultat)

        statut = resultat["Statut"]
        nb_entites = resultat["Nombre d'entités"]
        message = resultat["Message"]
        print(f"    {statut} — {nb_entites} entité(s)")
        if message:
            print(f"    {message}")
        print()

    # Rapport dans dossier_sortie (garanti d'exister).
    chemin_rapport = os.path.join(dossier_sortie, NOM_RAPPORT)
    try:
        creer_rapport(resultats, chemin_rapport)
        print("=" * 72)
        print("TRAITEMENT TERMINÉ")
        print(f"Réussites : {sum(r['Statut'] == 'SUCCÈS' for r in resultats)}")
        print(f"Fichiers vides : {sum(r['Statut'] == 'VIDE' for r in resultats)}")
        print(f"Échecs : {sum(r['Statut'] == 'ÉCHEC' for r in resultats)}")
        print(f"Rapport Excel : {chemin_rapport}")
    except Exception as exc:
        print(f"\nImpossible de créer le rapport Excel : {exc}")

    input("\nAppuyez sur Entrée pour fermer...")


if __name__ == "__main__":
    main()
